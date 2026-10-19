import json, sys, os
from openpyxl import load_workbook

path = r'e:/Arav Innovations/omnigrc/scratch/POC Criteria (2).xlsx'
if not os.path.exists(path):
    sys.exit('File not found')

wb = load_workbook(filename=path, read_only=True, data_only=True)
if 'Reports' not in wb.sheetnames:
    sys.exit('Reports sheet not found')
ws = wb['Reports']

# Assume first row is header
headers = []
rows = []
for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
    if i == 1:
        headers = [str(cell).strip() if cell is not None else '' for cell in row]
        continue
    row_dict = {}
    for header, cell in zip(headers, row):
        row_dict[header] = cell
    rows.append(row_dict)

print(json.dumps(rows, ensure_ascii=False, indent=2))
